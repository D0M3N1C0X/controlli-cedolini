"""
Il catalogo dei controlli: una scheda per controllo, come la knowledge base interna di un team paghe.

Quello che sta già nel codice si legge dal codice, così il catalogo non può restare indietro: nome, priorità,
cosa fare (controlli.py, build_rapporto.py), la formula del foglio (dal file Excel costruito), un esempio vero
(dalle segnalazioni del mese). A mano ci sono solo la regola in parole, le cause tipiche e i limiti.

    python src/build_catalogo.py
"""
import pandas as pd
from openpyxl import load_workbook

import build_rapporto
import config as C
import controlli

OUT = C.ROOT / "docs" / "catalogo-controlli.md"
FOGLIO = C.DELIVERABLES / "foglio_controllo_cedolini.xlsx"

# Per ogni controllo: dove sta nel foglio (colonne dei valori attesi e dell'esito) e le fonti nel registro.
FOGLIO_COLONNE = {
    "C01": ("Mese", ["X"], "AK"), "C02": ("Mese", ["Y", "Z", "AA"], "AL"), "C03": ("Mese", ["AB"], "AM"),
    "C04": ("Mese", ["AC"], "AN"), "C05": ("Mese", ["AD", "AE"], "AO"), "C06": ("Mese", ["AF", "AG"], "AP"),
    "C07": ("Mese", ["AH", "AI"], "AQ"), "C08": ("F24", ["C", "D", "F", "G"], "H"), "C09": ("Mese", ["AJ"], "AR"),
    "C10": ("Mese", [], "AS"), "C11": ("Mese", ["AW", "AX"], "AY"),
}
FONTI = {
    "C01": ["F01", "F02"], "C02": ["F03", "S04"], "C03": ["F05", "F10"], "C04": ["F06", "S03"], "C05": ["F04"],
    "C06": ["S01"], "C07": ["S02"], "C08": ["F09"], "C09": ["F08"], "C10": ["S06"], "C11": ["F11"],
}

SCHEDE = {
    "C01": {
        "regola": "Il tabellare è il minimo del livello in vigore (paga base, altri elementi, contingenza ed EDR) per la "
                  "percentuale di part-time, arrotondato al centesimo.",
        "cause": ["La tranche del CCNL non è stata caricata nel software paghe.",
                  "Livello sbagliato in anagrafica.",
                  "Part-time non riproporzionato, o riproporzionato due volte."],
        "limiti": ["Non controlla il superminimo né il suo assorbimento: quello è nel foglio Novembre 2026.",
                   "Vale per il CCNL Terziario Confcommercio: un altro contratto richiede la sua tabella."],
    },
    "C02": {
        "regola": "Uno scatto per ogni triennio dalla data di assunzione, in pagamento dal mese dopo il compimento, fino "
                  "a 10; l'importo è quello del livello per il numero di scatti e per il part-time.",
        "cause": ["Scatto maturato e non attivato nel software.",
                  "Data di assunzione sbagliata in anagrafica.",
                  "Scatti rimasti al livello precedente dopo un passaggio."],
        "limiti": ["Valorizza tutti gli scatti al livello attuale (S04): chi ha cambiato livello va controllato a mano.",
                   "Non considera anzianità convenzionali o periodi non utili alla maturazione."],
    },
    "C03": {
        "regola": "Il contributo IVS a carico del dipendente è l'imponibile arrotondato all'euro per l'aliquota "
                  "(9,19% se il cliente non ne indica un'altra), arrotondato al centesimo.",
        "cause": ["Aliquota della posizione INPS sbagliata, per esempio 9,49% dove non c'è CIGS.",
                  "Imponibile non arrotondato all'euro, o arrotondato due volte."],
        "limiti": ["Un'aliquota per tutto il cliente: FIS, CIGS e il contributo aggiuntivo dell'1% non sono modellati.",
                   "L'imponibile è il lordo: voci escluse dall'imponibile non sono gestite."],
    },
    "C04": {
        "regola": "Il rateo mensile di TFR è la retribuzione ordinaria (tabellare, scatti, superminimo) per 14 "
                  "mensilità, diviso 13,5 e diviso 12.",
        "cause": ["Base del TFR senza scatti o superminimo.",
                  "Divisore o numero di mensilità sbagliati nella configurazione."],
        "limiti": ["Non include voci variabili non occasionali né la rivalutazione del fondo.",
                   "Non sottrae il contributo dello 0,50% al Fondo pensioni lavoratori dipendenti."],
    },
    "C05": {
        "regola": "Il residuo ferie è il residuo del mese prima più un dodicesimo dei 26 giorni annui, meno le ferie "
                  "godute nel mese, a quattro decimali.",
        "cause": ["Maturazione del mese non registrata.",
                  "Ferie godute non scaricate dal residuo."],
        "limiti": ["Ferie in giorni, non in ore; ROL ed ex festività non sono gestiti.",
                   "Per chi è assunto nel mese il residuo di partenza è zero."],
    },
    "C06": {
        "regola": "Il lordo è la somma di tabellare, scatti, superminimo e straordinari; il netto è il lordo meno "
                  "contributi, IRPEF e altre trattenute.",
        "cause": ["Una voce aggiunta a mano e non sommata.",
                  "Una trattenuta inserita due volte."],
        "limiti": ["Una sola segnalazione per cedolino: se il lordo non torna, il netto non si controlla (D5).",
                   "Non ricalcola l'IRPEF: ne controlla solo l'effetto sul netto."],
    },
    "C07": {
        "regola": "Il netto non deve cambiare di oltre il 20%, in su o in giù, rispetto al mese prima, a meno che il "
                  "cedolino indichi un evento che lo spieghi.",
        "cause": ["Straordinari o arretrati pagati senza un evento dichiarato.",
                  "Voci duplicate, che gonfiano il netto."],
        "limiti": ["La soglia del 20% va tarata sui primi mesi di uso reale (S02).",
                   "Un evento dichiarato spegne il controllo, anche se l'importo è sbagliato."],
    },
    "C08": {
        "regola": "Per ogni cliente, le ritenute versate con il codice tributo 1001 e i contributi a carico dei "
                  "dipendenti nell'F24 sono uguali alla somma dei cedolini del mese.",
        "cause": ["Un dipendente rimasto fuori dalla delega.",
                  "Delega preparata prima di una correzione ai cedolini."],
        "limiti": ["La sezione INPS è semplificata alla sola quota dei dipendenti (D12).",
                   "Le addizionali regionali e comunali (codici 3802 e 3848) non sono riconciliate."],
    },
    "C09": {
        "regola": "Il carattere di controllo del codice fiscale (il sedicesimo) è quello che il DM 23 dicembre 1976 "
                  "ricava dai primi quindici.",
        "cause": ["Codice trascritto a mano con un carattere sbagliato."],
        "limiti": ["Controlla solo il carattere di controllo, non che il codice corrisponda alla persona.",
                   "I codici con omocodia sono validi e passano, come devono."],
    },
    "C10": {
        "regola": "Livello, part-time, superminimo e data di assunzione sono uguali al mese prima, a meno che il "
                  "cedolino indichi un evento.",
        "cause": ["Anagrafica modificata per errore.",
                  "Un cambio vero non comunicato come evento."],
        "limiti": ["Un errore presente fin dall'assunzione non cambia da un mese all'altro e resta invisibile: è il "
                   "limite dichiarato e tenuto vero da un test.",
                   "Chi non ha un cedolino il mese prima non viene confrontato."],
    },
    "C11": {
        "regola": "I dodicesimi di tredicesima (da gennaio) e di quattordicesima (da luglio) sono i mesi di servizio "
                  "nel periodo, fino al mese del cedolino; il mese di assunzione conta se ha almeno 15 giorni.",
        "cause": ["Mese di assunzione con meno di 15 giorni contato intero.",
                  "Quattordicesima contata da gennaio invece che da luglio.",
                  "Un dodicesimo perso senza assenze che lo spieghino."],
        "limiti": ["Controlla il numero di dodicesimi, non l'importo (D14).",
                   "Non considera i periodi non retribuiti che riducono i ratei, come l'aspettativa.",
                   "Le colonne dei ratei sono facoltative: senza, C11 non si applica (D15)."],
    },
}


def _numero(v) -> str:
    try:
        x = float(v)
    except (TypeError, ValueError):
        return str(v)
    return build_rapporto.euro(x)


def _giorni(v) -> str:
    return f"{float(v):.4f}".replace(".", ",")


def _esempio(codice: str, an: pd.DataFrame) -> str:
    righe = an[an["controllo"] == codice]
    if righe.empty:
        return "Nessuna segnalazione nel mese di esempio."
    r = righe.iloc[0]
    chi = f"Dipendente {r['matricola']}, cliente {r['cliente']}"
    if codice == "C05":
        return f"{chi}: residuo atteso {_giorni(r['atteso'])} giorni, sul cedolino {_giorni(r['trovato'])}."
    if codice == "C07":
        variazione = float(r["trovato"]) / float(r["atteso"]) - 1
        segno = "+" if variazione > 0 else ""
        return (f"{chi}: netto di agosto {_numero(r['atteso'])}, di settembre {_numero(r['trovato'])} "
                f"({segno}{build_rapporto.pct(variazione)}), nessun evento sul cedolino.")
    if codice == "C08":
        return (f"Cliente {r['cliente']}: somma dei cedolini {_numero(r['atteso'])}, nell'F24 {_numero(r['trovato'])}, "
                f"differenza {_numero(r['differenza'])}.")
    if codice == "C09":
        return f"{chi}: il codice «{r['trovato']}» ha un carattere di controllo non valido."
    if codice in ("C10", "C11"):
        return f"{chi}: atteso «{r['atteso']}», sul cedolino «{r['trovato']}»."
    return (f"{chi}: atteso {_numero(r['atteso'])}, sul cedolino {_numero(r['trovato'])}, "
            f"differenza {_numero(r['differenza'])}.")


def scheda(codice: str, wb, an: pd.DataFrame) -> list[str]:
    nome, descrizione, priorita = controlli.CONTROLLI[codice]
    s = SCHEDE[codice]
    foglio, attese, esito = FOGLIO_COLONNE[codice]
    ws = wb[foglio]
    righe = [f"## {codice} · {nome}", "",
             f"**Priorità {priorita}.** {descrizione}.", "",
             f"**La regola.** {s['regola']}", "",
             "**Fonti.** " + ", ".join(f"[{f}](fonti.md)" for f in FONTI[codice]) + ".", "",
             f"**Nel foglio** (foglio {foglio}, riga 5; le altre righe sono uguali):", "", "```"]
    for col in [*attese, esito]:
        righe.append(f"{col}  {ws[f'{col}4'].value}")
        righe.append(f"    {ws[f'{col}5'].value}")
    righe += ["```", "",
              f"**Un esempio dal mese di settembre 2026.** {_esempio(codice, an)}", "",
              "**Cause tipiche.**", "", *[f"- {c}" for c in s["cause"]], "",
              f"**Cosa fare.** {build_rapporto.COSA_FARE[codice]}.", "",
              "**Limiti.**", "", *[f"- {x}" for x in s["limiti"]], ""]
    return righe


def testo() -> str:
    wb = load_workbook(FOGLIO)
    an = pd.read_csv(C.REPORT / "anomalie.csv", dtype={"matricola": str}, keep_default_na=False)
    righe = ["# Catalogo dei controlli", "",
             "Una scheda per controllo: cosa controlla, da quale regola, dove sta nel foglio, un esempio vero e cosa "
             "fare. Generato da `src/build_catalogo.py`: nome, priorità, formule ed esempi vengono dal codice e dal "
             "foglio, quindi non restano indietro.", "",
             "| Controllo | Priorità | Cosa controlla |", "|---|---|---|"]
    for codice, (nome, descrizione, priorita) in controlli.CONTROLLI.items():
        righe.append(f"| [{codice} · {nome}](#{codice.lower()}--{nome.lower().replace(' ', '-')}) | {priorita} | "
                     f"{descrizione} |")
    righe.append("")
    for codice in controlli.CONTROLLI:
        righe += scheda(codice, wb, an)
    return "\n".join(righe).rstrip("\n") + "\n"


def write() -> None:
    OUT.write_text(testo(), encoding="utf-8")


if __name__ == "__main__":
    write()
    print(f"catalogo -> {OUT.relative_to(C.ROOT)}")
