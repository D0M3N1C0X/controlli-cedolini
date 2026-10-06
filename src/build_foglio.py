"""
Scrive deliverables/foglio_controllo_cedolini.xlsx: i controlli come formule, per chi non usa
Python. Si apre in Excel, in LibreOffice e in Google Sheets (File > Importa): si incollano i
cedolini del mese nei fogli Mese e Mese precedente e gli esiti si aggiornano da soli.

Convenzioni: testo blu = dato da incollare o parametro, nero = formula, verde = collegamento a un
altro foglio, giallo = parametro che cambia i risultati. Nessuna formula a matrice dinamica.
Il foglio Riconciliazione confronta ogni formula con il risultato di Python.
"""
import math
import string
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.hyperlink import Hyperlink

import codice_fiscale
import config as C
import controlli
import regole as R
from deterministic import normalise

OUTPUT = C.DELIVERABLES / "foglio_controllo_cedolini.xlsx"
RIGHE = 250          # righe con le formule già pronte nei fogli Mese, Mese precedente e Novembre 2026
REPO = "https://github.com/D0M3N1C0X/controlli-cedolini"

FONT = "Arial"
BLUE, GREEN, INK, MUTED, WHITE = "0000FF", "008000", "1B2430", "5F6B7A", "FFFFFF"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
YELLOW = PatternFill("solid", fgColor="FFFF00")
RED_FILL = PatternFill("solid", fgColor="F5C6C2")
GREEN_FILL = PatternFill("solid", fgColor="D5ECDC")
EUR = '"€"#,##0.00;-"€"#,##0.00;"-"'
NUM2, NUM4, PCT = "#,##0.00", "0.0000", "0.00%"


def font(color=INK, bold=False, italic=False, size=10):
    return Font(name=FONT, color=color, bold=bold, italic=italic, size=size)


def put(ws, ref, value, *, color=INK, bold=False, italic=False, size=10, fmt=None, fill=None, wrap=False):
    cell = ws[ref]
    cell.value = value
    cell.font = font(color, bold, italic, size)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return cell


def header(ws, row, labels, start=1, height=32):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=label)
        c.font = font(WHITE, bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = height


def title(ws, text, subtitle=None):
    put(ws, "A1", text, bold=True, size=14)
    if subtitle:
        put(ws, "A2", subtitle, color=MUTED, italic=True)


def widths(ws, spec):
    for k, v in spec.items():
        ws.column_dimensions[k].width = v


def plain(x):
    if hasattr(x, "item"):
        x = x.item()
    if isinstance(x, float) and math.isnan(x):
        return None
    return x


def esito_rule(ws, ref):
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"ERRORE"'], fill=RED_FILL))
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"OK"'], fill=GREEN_FILL))


class Foglio:
    def __init__(self, o: dict):
        self.o = o
        self.wb = Workbook()
        self.R = {}
        self.checks = []

    def check(self, area, item, value, ref):
        self.checks.append((area, item, plain(value), ref))

    def build(self, path: Path):
        names = ["Leggimi", "Riepilogo", "Mese", "Mese precedente", "F24", "Novembre 2026", "CCNL", "Parametri",
                 "Codice fiscale", "Riconciliazione"]
        self.wb.active.title = names[0]
        for n in names[1:]:
            self.wb.create_sheet(n)
        self.parametri(self.wb["Parametri"])
        self.ccnl(self.wb["CCNL"])
        self.tabelle_cf(self.wb["Codice fiscale"])
        self.agosto(self.wb["Mese precedente"])
        self.settembre(self.wb["Mese"])
        self.f24(self.wb["F24"])
        self.novembre(self.wb["Novembre 2026"])
        self.riepilogo(self.wb["Riepilogo"])
        self.riconciliazione(self.wb["Riconciliazione"])
        self.leggimi(self.wb["Leggimi"], names)
        for ws in self.wb.worksheets:
            ws.sheet_view.showGridLines = False
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.properties.creator = "Domenico Perroni"
        self.wb.properties.title = "Controlli cedolini"
        self.wb.properties.created = self.wb.properties.modified = datetime(2026, 1, 1)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(path)
        normalise(path)

    # ---- Inputs -------------------------------------------------------------------------------
    def parametri(self, ws):
        title(ws, "Parametri", "Le regole dei controlli. I valori gialli cambiano gli esiti: vanno presi dal cliente.")
        widths(ws, {"A": 3, "B": 44, "C": 14, "D": 90})
        rows = [
            ("mese", "Mese controllato (primo giorno)", date.fromisoformat(C.MESE + "-01"), "DD/MM/YYYY", None,
             "Il mese dei cedolini incollati nel foglio Mese: conta per gli scatti."),
            ("aliquota", "Contributo IVS a carico del dipendente", C.ALIQUOTA_DIPENDENTE, PCT, YELLOW,
             "9,19% generale. Va preso dalla posizione INPS del cliente: FIS, CIGS e l'1% sopra la prima fascia pensionabile non sono modellati."),
            ("divisore_tfr", "Divisore TFR", C.DIVISORE_TFR, "0.0", None, "Art. 2120 c.c.: quota annua = retribuzione annua / 13,5."),
            ("mensilita", "Mensilità", C.MENSILITA, "0", None, "CCNL Terziario: tredicesima e quattordicesima."),
            ("ferie", "Ferie annue, giorni lavorativi", C.FERIE_ANNUE, "0", None, "CCNL Terziario: 26 giorni, maturati per dodicesimi."),
            ("soglia", "Variazione del netto da spiegare", C.SOGLIA_VARIAZIONE, PCT, YELLOW,
             "Oltre questa soglia sul mese prima, senza un evento, il cedolino va guardato."),
            ("tolleranza", "Tolleranza, euro", C.TOLLERANZA, "0.00", None, "Un centesimo di arrotondamento."),
            ("tolleranza_ferie", "Tolleranza, giorni di ferie", C.TOLLERANZA_FERIE, "0.00", None, ""),
            ("scatti_max", "Scatti massimi", C.SCATTI_MAX, "0", None, "CCNL Terziario: 10 scatti triennali."),
        ]
        header(ws, 4, ["", "Parametro", "Valore", "Fonte e note"], height=20)
        for r, (key, label, value, fmt, fill, note) in enumerate(rows, start=5):
            put(ws, f"B{r}", label)
            put(ws, f"C{r}", value, color=BLUE, fmt=fmt, fill=fill)
            put(ws, f"D{r}", note, color=MUTED, italic=True)
            self.R[key] = f"Parametri!$C${r}"

    def ccnl(self, ws):
        t = R.tabella()
        title(ws, "CCNL Terziario Distribuzione e Servizi (Confcommercio): minimi mensili e scatti",
              "Accordo di rinnovo del 22 marzo 2024. Minimo = paga base + contingenza (+ 5,16 € al 7° livello). Fonti in docs/fonti.md.")
        widths(ws, {"A": 10, **{col(k): 14 for k in range(2, 10)}})
        header(ws, 4, ["Livello", "Paga base", "Contingenza", "Altri elementi", "Minimo dal 03/2025",
                       "Minimo dal 11/2025", "Minimo dal 11/2026", "Minimo dal 02/2027", "Scatto"])
        first = 5
        for i, (lv, r) in enumerate(t.iterrows(), start=first):
            put(ws, f"A{i}", lv, color=BLUE)
            for j, c in enumerate(["paga_base", "contingenza", "altri_elementi", "minimo_2025_03", "minimo_2025_11",
                                   "minimo_2026_11", "minimo_2027_02", "scatto"]):
                put(ws, f"{col(2 + j)}{i}", float(r[c]), color=BLUE, fmt=NUM2)
        last = first + len(t) - 1
        rng = lambda letter: f"CCNL!${letter}${first}:${letter}${last}"
        self.R.update({"lv": rng("A"), "min": rng("F"), "min_next": rng("G"), "scatto": rng("I")})

    def tabelle_cf(self, ws):
        title(ws, "Codice fiscale: valori per il carattere di controllo", "DM 23 dicembre 1976: posizioni dispari e pari.")
        widths(ws, {"A": 10, "B": 12, "C": 12})
        header(ws, 4, ["Carattere", "Dispari", "Pari"], height=20)
        chars = "0123456789" + string.ascii_uppercase
        for i, c in enumerate(chars, start=5):
            put(ws, f"A{i}", c, color=BLUE)
            put(ws, f"B{i}", codice_fiscale.DISPARI[c], color=BLUE)
            put(ws, f"C{i}", codice_fiscale.PARI[c], color=BLUE)
        last = 4 + len(chars)
        self.R.update({"cf_c": f"'Codice fiscale'!$A$5:$A${last}", "cf_d": f"'Codice fiscale'!$B$5:$B${last}",
                       "cf_p": f"'Codice fiscale'!$C$5:$C${last}"})

    def agosto(self, ws):
        p = self.o["precedente"]
        title(ws, "Mese precedente: i cedolini del mese prima", "Servono il netto e il residuo ferie, per i controlli C05 e C07. Nei dati di esempio: agosto 2026.")
        widths(ws, {"A": 12, "B": 9, "C": 14, "D": 14})
        header(ws, 4, ["Matricola", "Cliente", "Netto", "Residuo ferie"], height=20)
        first = 5
        for i, t in enumerate(p.itertuples(index=False), start=first):
            put(ws, f"A{i}", t.matricola, color=BLUE)
            put(ws, f"B{i}", t.cliente, color=BLUE)
            put(ws, f"C{i}", float(t.netto), color=BLUE, fmt=NUM2)
            put(ws, f"D{i}", float(t.ferie_residuo), color=BLUE, fmt=NUM4)
        last = first + RIGHE - 1                 # spazio per incollare più righe
        self.R.update({"ag_m": f"'Mese precedente'!$A${first}:$A${last}", "ag_n": f"'Mese precedente'!$C${first}:$C${last}",
                       "ag_f": f"'Mese precedente'!$D${first}:$D${last}"})
        ws.freeze_panes = "A5"

    # ---- The checks ------------------------------------------------------------------------------
    def settembre(self, ws):
        R_, m, a = self.R, self.o["mese"], self.o["attesi"]
        title(ws, "Mese: cedolini e controlli (nei dati di esempio, settembre 2026)",
              f"Colonne A-W: i dati del software paghe (blu), fino a {RIGHE} righe. Colonne X-AS: valori attesi ed esiti, in formula.")
        inputs = ["matricola", "cliente", "codice_fiscale", "livello", "data_assunzione", "part_time", "evento",
                  "superminimo_assorbibile", "tabellare", "scatti_n", "scatti", "superminimo", "straordinari", "lordo",
                  "contributi_inps", "imponibile_irpef", "irpef", "altre_trattenute", "netto", "quota_tfr",
                  "ferie_maturate", "ferie_godute", "ferie_residuo"]
        heads = [s.replace("_", " ").capitalize() for s in inputs] + [
            "Tabellare atteso", "Mesi di servizio", "Scatti attesi", "Importo scatti atteso", "Contributi attesi",
            "TFR atteso", "Residuo mese prima", "Residuo atteso", "Lordo atteso", "Netto atteso", "Netto mese prima",
            "Variazione", "Controllo CF",
            "C01 Minimo", "C02 Scatti", "C03 INPS", "C04 TFR", "C05 Ferie", "C06 Quadratura", "C07 Variazione",
            "C09 Cod. fiscale", "Errori"]
        header(ws, 4, heads, height=44)
        widths(ws, {**{col(k): 11 for k in range(1, len(heads) + 1)}, "C": 19, "G": 14})
        first = 5
        assert len(m) <= RIGHE, "più cedolini delle righe preparate nel foglio"
        last = first + RIGHE - 1
        self.rows = (first, last)
        for i, t in enumerate(m.itertuples(index=False), start=first):
            for j, name in enumerate(inputs):
                v = getattr(t, name)
                if name == "data_assunzione":
                    v = date.fromisoformat(v)
                elif name == "superminimo_assorbibile":
                    v = str(v).lower() == "true"
                elif isinstance(v, float) or name in ("part_time",):
                    v = float(v)
                elif name == "scatti_n":
                    v = int(v)
                put(ws, f"{col(1 + j)}{i}", plain(v), color=BLUE,
                    fmt="DD/MM/YYYY" if name == "data_assunzione" else (NUM4 if "ferie" in name else
                                                                         (NUM2 if isinstance(v, float) else None)))
        # le formule coprono RIGHE righe, così si possono incollare i cedolini di un cliente più grande;
        # le righe vuote restano vuote
        for i in range(first, last + 1):
            g = lambda expr: f'=IF($A{i}="","",{expr})'
            lookup = lambda rng: f"INDEX({rng},MATCH(D{i},{R_['lv']},0))"
            put(ws, f"X{i}", g(f"ROUND({lookup(R_['min'])}*F{i},2)"), fmt=NUM2)
            put(ws, f"Y{i}", g(f"(YEAR({R_['mese']})*12+MONTH({R_['mese']}))-(YEAR(E{i})*12+MONTH(E{i}))"))
            put(ws, f"Z{i}", g(f"MIN({R_['scatti_max']},MAX(0,INT((Y{i}-1)/36)))"))
            put(ws, f"AA{i}", g(f"ROUND({lookup(R_['scatto'])}*Z{i}*F{i},2)"), fmt=NUM2)
            put(ws, f"AB{i}", g(f"ROUND(N{i}*{R_['aliquota']},2)"), fmt=NUM2)
            put(ws, f"AC{i}", g(f"ROUND((I{i}+K{i}+L{i})*{R_['mensilita']}/{R_['divisore_tfr']}/12,2)"), fmt=NUM2)
            put(ws, f"AD{i}", g(f'IFERROR(INDEX({R_["ag_f"]},MATCH(A{i},{R_["ag_m"]},0)),0)'), color=GREEN, fmt=NUM4)
            put(ws, f"AE{i}", g(f"ROUND(AD{i}+ROUND({R_['ferie']}/12,4)-V{i},4)"), fmt=NUM4)
            put(ws, f"AF{i}", g(f"ROUND(I{i}+K{i}+L{i}+M{i},2)"), fmt=NUM2)
            put(ws, f"AG{i}", g(f"ROUND(N{i}-O{i}-Q{i}-R{i},2)"), fmt=NUM2)
            put(ws, f"AH{i}", g(f'IFERROR(INDEX({R_["ag_n"]},MATCH(A{i},{R_["ag_m"]},0)),"")'), color=GREEN, fmt=NUM2)
            put(ws, f"AI{i}", g(f'IF(AH{i}="","",S{i}/AH{i}-1)'), fmt=PCT)
            odd = "{1,3,5,7,9,11,13,15}"
            even = "{2,4,6,8,10,12,14}"
            put(ws, f"AJ{i}", g(f"CHAR(65+MOD(SUMPRODUCT(SUMIF({R_['cf_c']},MID(C{i},{odd},1),{R_['cf_d']}))"
                                f"+SUMPRODUCT(SUMIF({R_['cf_c']},MID(C{i},{even},1),{R_['cf_p']})),26))"))
            tol, tolf = R_["tolleranza"], R_["tolleranza_ferie"]
            put(ws, f"AK{i}", g(f'IF(ABS(I{i}-X{i})>{tol},"ERRORE","OK")'))
            put(ws, f"AL{i}", g(f'IF(ABS(K{i}-AA{i})>{tol},"ERRORE","OK")'))
            put(ws, f"AM{i}", g(f'IF(ABS(O{i}-AB{i})>{tol},"ERRORE","OK")'))
            put(ws, f"AN{i}", g(f'IF(ABS(T{i}-AC{i})>{tol},"ERRORE","OK")'))
            put(ws, f"AO{i}", g(f'IF(ABS(W{i}-AE{i})>{tolf},"ERRORE","OK")'))
            put(ws, f"AP{i}", g(f'IF(OR(ABS(N{i}-AF{i})>{tol},ABS(S{i}-AG{i})>{tol}),"ERRORE","OK")'))
            put(ws, f"AQ{i}", g(f'IF(AI{i}="","OK",IF(AND(ABS(AI{i})>{R_["soglia"]},G{i}=""),"ERRORE","OK"))'))
            put(ws, f"AR{i}", g(f'IF(AND(LEN(C{i})=16,RIGHT(C{i},1)=AJ{i}),"OK","ERRORE")'))
            put(ws, f"AS{i}", g(f'COUNTIF(AK{i}:AR{i},"ERRORE")'), bold=True)
        for i, t in enumerate(m.itertuples(index=False), start=first):
            k = i - first
            if k % 6 == 0:
                for letter, name in (("X", "tabellare"), ("Z", "scatti_n"), ("AA", "scatti"), ("AB", "contributi_inps"),
                                     ("AC", "quota_tfr"), ("AE", "ferie_residuo"), ("AG", "netto")):
                    self.check("Mese", f"{t.matricola} {name} atteso", a.loc[k, name], f"Mese!{letter}{i}")
        esito_rule(ws, f"AK{first}:AR{last}")
        E = lambda letter: f"Mese!${letter}${first}:${letter}${last}"
        self.R.update({"s_cli": E("B"), "s_m": E("A"), "s_irpef": E("Q"), "s_inps": E("O"), "s_err": E("AS"),
                       **{f"s_{c}": E(letter) for c, letter in zip(
                           ["C01", "C02", "C03", "C04", "C05", "C06", "C07", "C09"],
                           ["AK", "AL", "AM", "AN", "AO", "AP", "AQ", "AR"])}})
        # each Python finding must be an ERRORE in the sheet, and every other row OK
        an = self.o["anomalie"]
        colonna = {"C01": "AK", "C02": "AL", "C03": "AM", "C04": "AN", "C05": "AO", "C06": "AP", "C07": "AQ", "C09": "AR"}
        trovate = set(zip(an["controllo"], an["matricola"]))
        for k, t in enumerate(m.itertuples(index=False)):
            for codice, letter in colonna.items():
                if (codice, t.matricola) in trovate or k % 10 == 0:
                    esito = "ERRORE" if (codice, t.matricola) in trovate else "OK"
                    self.check("Mese", f"{t.matricola} {codice}", esito, f"Mese!{letter}{first + k}")
        ws.freeze_panes = "D5"
        ws.auto_filter.ref = f"A4:AS{last}"

    def f24(self, ws):
        R_, rec = self.R, self.o["riconciliazione_f24"]
        title(ws, "F24 del mese: deleghe contro cedolini",
              "Ritenute IRPEF (codice tributo 1001) e contributi a carico dei dipendenti, per cliente.")
        widths(ws, {"A": 10, **{col(k): 16 for k in range(2, 10)}})
        header(ws, 4, ["Cliente", "Ritenute in F24", "Ritenute nei cedolini", "Differenza", "Contributi in F24",
                       "Contributi nei cedolini", "Differenza", "C08 Esito"])
        for i, t in enumerate(rec.itertuples(index=False), start=5):
            put(ws, f"A{i}", t.cliente, color=BLUE)
            put(ws, f"B{i}", float(t.ritenute_1001), color=BLUE, fmt=NUM2)
            put(ws, f"C{i}", f'=ROUND(SUMIFS({R_["s_irpef"]},{R_["s_cli"]},A{i}),2)', fmt=NUM2)
            put(ws, f"D{i}", f"=ROUND(B{i}-C{i},2)", fmt=NUM2)
            put(ws, f"E{i}", float(t.contributi_dipendente), color=BLUE, fmt=NUM2)
            put(ws, f"F{i}", f'=ROUND(SUMIFS({R_["s_inps"]},{R_["s_cli"]},A{i}),2)', fmt=NUM2)
            put(ws, f"G{i}", f"=ROUND(E{i}-F{i},2)", fmt=NUM2)
            put(ws, f"H{i}", f'=IF(OR(ABS(D{i})>{R_["tolleranza"]},ABS(G{i})>{R_["tolleranza"]}),"ERRORE","OK")')
            self.check("F24", f"Cliente {t.cliente} differenza ritenute", float(t.diff_ritenute), f"F24!D{i}")
            self.check("F24", f"Cliente {t.cliente} differenza contributi", float(t.diff_contributi), f"F24!G{i}")
        last = 4 + len(rec)
        esito_rule(ws, f"H5:H{last}")
        self.R.update({"f24_cli": f"F24!$A$5:$A${last}", "f24_es": f"F24!$H$5:$H${last}"})

    def novembre(self, ws):
        R_, n = self.R, self.o["novembre"]
        title(ws, "Novembre 2026: la nuova tranche del CCNL Terziario",
              "Si aggiorna dal foglio Mese. Superminimo assorbibile: verificare la lettera di assunzione di ciascuno.")
        widths(ws, {"A": 12, "B": 9, "C": 9, "D": 10, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12, "J": 14, "K": 14})
        header(ws, 4, ["Matricola", "Cliente", "Livello", "Part-time", "Superminimo", "Assorbibile",
                       "Aumento mensile", "Assorbito", "Da pagare", "Nuovo tabellare", "Nuovo superminimo"])
        first = 5
        for i in range(first, first + RIGHE):
            g = lambda expr: f'=IF($A{i}="","",{expr})'
            for letter, src in (("A", "A"), ("B", "B"), ("C", "D"), ("D", "F"), ("E", "L"), ("F", "H")):
                fmt = NUM2 if letter == "E" else None
                put(ws, f"{letter}{i}", f'=IF(Mese!$A{i}="","",Mese!{src}{i})', color=GREEN, fmt=fmt)
            lk = lambda rng: f"INDEX({rng},MATCH(C{i},{R_['lv']},0))"
            put(ws, f"G{i}", g(f"ROUND(({lk(R_['min_next'])}-{lk(R_['min'])})*D{i},2)"), fmt=NUM2)
            put(ws, f"H{i}", g(f"IF(F{i},MIN(E{i},G{i}),0)"), fmt=NUM2)
            put(ws, f"I{i}", g(f"ROUND(G{i}-H{i},2)"), fmt=NUM2)
            put(ws, f"J{i}", g(f"ROUND({lk(R_['min_next'])}*D{i},2)"), fmt=NUM2)
            put(ws, f"K{i}", g(f"ROUND(E{i}-H{i},2)"), fmt=NUM2)
        for i, t in enumerate(n.itertuples(index=False), start=first):
            if (i - first) % 5 == 0:
                for letter, name in (("G", "aumento"), ("H", "assorbito"), ("I", "da_pagare"), ("J", "nuovo_tabellare")):
                    self.check("Novembre 2026", f"{t.matricola} {name}", float(getattr(t, name)), f"'Novembre 2026'!{letter}{i}")
        last = first + RIGHE - 1
        N = lambda letter: f"'Novembre 2026'!${letter}${first}:${letter}${last}"
        self.R.update({"n_cli": N("B"), "n_aum": N("G"), "n_ass": N("H"), "n_pag": N("I")})
        ws.freeze_panes = "B5"

    def riepilogo(self, ws):
        R_, o = self.R, self.o
        title(ws, "Riepilogo del mese", "Errori per controllo e cliente, e il costo della tranche di novembre.")
        clienti = list(C.CLIENTI)
        widths(ws, {"A": 8, "B": 40, **{col(k): 13 for k in range(3, 4 + len(clienti))}})
        header(ws, 4, ["Codice", "Controllo", *[f"Cliente {c}" for c in clienti], "Totale"])
        an = o["anomalie"]
        r = 5
        for codice, (voce, desc, _) in controlli.CONTROLLI.items():
            put(ws, f"A{r}", codice, bold=True)
            put(ws, f"B{r}", voce)
            for j, c in enumerate(clienti):
                letter = col(3 + j)
                if codice == "C08":
                    f = f'=COUNTIFS({R_["f24_cli"]},"{c}",{R_["f24_es"]},"ERRORE")'
                else:
                    f = f'=COUNTIFS({R_["s_cli"]},"{c}",{R_["s_" + codice]},"ERRORE")'
                put(ws, f"{letter}{r}", f)
                self.check("Riepilogo", f"{codice} cliente {c}", int(((an["controllo"] == codice) & (an["cliente"] == c)).sum()),
                           f"Riepilogo!{letter}{r}")
            put(ws, f"{col(3 + len(clienti))}{r}", f"=SUM(C{r}:{col(2 + len(clienti))}{r})", bold=True)
            r += 1
        put(ws, f"B{r}", "Segnalazioni", bold=True)
        for j in range(len(clienti) + 1):
            letter = col(3 + j)
            put(ws, f"{letter}{r}", f"=SUM({letter}5:{letter}{r - 1})", bold=True)
        self.check("Riepilogo", "Segnalazioni totali", len(an), f"Riepilogo!{col(3 + len(clienti))}{r}")
        for k in range(len(clienti) + 1):
            ws.conditional_formatting.add(f"{col(3 + k)}5:{col(3 + k)}{r - 1}",
                                          CellIsRule(operator="greaterThan", formula=["0"], fill=RED_FILL))
        r += 3
        put(ws, f"A{r}", "Tranche del 1° novembre 2026, al mese", bold=True, size=11)
        r += 1
        header(ws, r, ["", "Voce", *[f"Cliente {c}" for c in clienti], "Totale"], height=20)
        rn = o["riepilogo_novembre"].set_index("cliente")
        for label, key, f in (("Aumento dei minimi", "aumento", R_["n_aum"]), ("Assorbito dai superminimi", "assorbito", R_["n_ass"]),
                              ("Da pagare in più ogni mese", "da_pagare", R_["n_pag"])):
            r += 1
            put(ws, f"B{r}", label, bold=key == "da_pagare")
            for j, c in enumerate(clienti):
                letter = col(3 + j)
                put(ws, f"{letter}{r}", f'=SUMIFS({f},{R_["n_cli"]},"{c}")', fmt=EUR, bold=key == "da_pagare")
                self.check("Riepilogo", f"Novembre {key} cliente {c}", float(rn.loc[c, key]), f"Riepilogo!{letter}{r}")
            put(ws, f"{col(3 + len(clienti))}{r}", f"=SUM(C{r}:{col(2 + len(clienti))}{r})", fmt=EUR, bold=True)
        r += 1
        put(ws, f"B{r}", f"Costo annuo lordo ({C.MENSILITA} mensilità)", bold=True)
        for j in range(len(clienti) + 1):
            letter = col(3 + j)
            put(ws, f"{letter}{r}", f"={letter}{r - 1}*{R_['mensilita']}", fmt=EUR, bold=True)
        put(ws, f"B{r + 1}", "Lordo dipendente, senza contributi a carico dell'azienda.", color=MUTED, italic=True)

    def riconciliazione(self, ws):
        title(ws, "Riconciliazione: formule del foglio contro Python",
              "La colonna D è scritta da src/build_foglio.py con i risultati di pandas; la colonna E è la formula.")
        widths(ws, {"A": 16, "B": 52, "C": 3, "D": 18, "E": 18, "F": 14, "G": 8})
        header(ws, 6, ["Foglio", "Voce", "", "Python", "Foglio", "Differenza", "Uguale"])
        first, last = 7, 6 + len(self.checks)
        put(ws, "B3", (f'=IF(COUNTIF(G{first}:G{last},"No")=0,"Tutti i "&COUNTA(B{first}:B{last})&" controlli tornano",'
                       f'COUNTIF(G{first}:G{last},"No")&" controlli su "&COUNTA(B{first}:B{last})&" non tornano")'),
            bold=True, size=12)
        for r, (area, item, value, ref) in enumerate(self.checks, start=first):
            put(ws, f"A{r}", area)
            put(ws, f"B{r}", item)
            fmt = "0.000000" if isinstance(value, float) else None
            put(ws, f"D{r}", value, color=BLUE, fmt=fmt)
            put(ws, f"E{r}", f"={ref}", color=GREEN, fmt=fmt)
            put(ws, f"F{r}", f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),E{r}-D{r},"")', fmt="0.0E+00")
            put(ws, f"G{r}", (f'=IF(AND(D{r}="",E{r}=""),"Sì",IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),'
                              f'IF(ABS(E{r}-D{r})<=1E-9*MAX(1,ABS(D{r})),"Sì","No"),IF(D{r}=E{r},"Sì","No")))'))
        for text, fill in (("Sì", GREEN_FILL), ("No", RED_FILL)):
            ws.conditional_formatting.add(f"G{first}:G{last}", CellIsRule(operator="equal", formula=[f'"{text}"'], fill=fill))
        ws.freeze_panes = "A7"
        self.recon = "Riconciliazione!B3"

    def leggimi(self, ws, names):
        widths(ws, {"A": 3, "B": 22, "C": 100})
        put(ws, "B2", "Controlli cedolini", bold=True, size=18)
        put(ws, "B3", "Nove controlli sui cedolini del mese, prima dell'invio al cliente: CCNL Terziario Confcommercio",
            color=MUTED, size=12)
        put(ws, "B5", "Come si usa", bold=True, size=11)
        for i, text in enumerate([
            f"1. Imposta il mese nei Parametri. Incolla i cedolini del mese nel foglio Mese (colonne A-W, fino a {RIGHE} righe) e quelli del mese prima nel foglio Mese precedente.",
            "2. Incolla gli importi delle deleghe F24 nel foglio F24 (colonne B ed E).",
            "3. Controlla i parametri gialli nel foglio Parametri: l'aliquota INPS va presa dal cliente.",
            "4. Leggi il Riepilogo e filtra il foglio Mese sulla colonna Errori maggiore di zero.",
        ], start=6):
            put(ws, f"C{i}", text)
        put(ws, "B11", "I controlli", bold=True, size=11)
        for i, (codice, (voce, desc, prio)) in enumerate(controlli.CONTROLLI.items(), start=12):
            put(ws, f"B{i}", f"{codice} {voce}")
            put(ws, f"C{i}", f"{desc}. Priorità {prio}.")
        r = 12 + len(controlli.CONTROLLI) + 1
        put(ws, f"B{r}", "Fogli", bold=True, size=11)
        purpose = {
            "Riepilogo": "Errori per controllo e cliente; costo della tranche di novembre.",
            "Mese": "I cedolini del mese con valori attesi ed esiti, riga per riga.",
            "Mese precedente": "Netto e residuo ferie del mese prima.",
            "F24": "Deleghe contro somma dei cedolini, per cliente.",
            "Novembre 2026": "Aumento dei minimi dal 1° novembre per dipendente, con l'assorbimento del superminimo.",
            "CCNL": "Minimi e scatti del CCNL Terziario.",
            "Parametri": "Aliquota, divisore TFR, ferie, soglie.",
            "Codice fiscale": "Valori per il carattere di controllo.",
            "Riconciliazione": "Ogni formula confrontata con Python.",
        }
        for i, n in enumerate(names[1:], start=r + 1):
            cell = put(ws, f"B{i}", n, color="1F5FA8")
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{n}'!A1")
            put(ws, f"C{i}", purpose[n])
        r = r + len(names) + 1
        put(ws, f"B{r}", "Verifica", bold=True, size=11)
        put(ws, f"C{r}", f"={self.recon}", color=GREEN, bold=True)
        put(ws, f"B{r + 2}", "Dati", bold=True, size=11)
        put(ws, f"C{r + 2}", "Sintetici: tre clienti e 122 dipendenti inventati, con errori inseriti apposta. I minimi "
                             "del CCNL sono reali (fonti in docs/fonti.md). IRPEF non calcolata: serve solo a riconciliare l'F24.",
            wrap=True)
        ws.row_dimensions[r + 2].height = 28
        put(ws, f"B{r + 3}", "Limiti", bold=True, size=11)
        put(ws, f"C{r + 3}", "Non trova un errore coerente: un livello sbagliato in anagrafica dà un cedolino che torna. "
                             "Non è consulenza del lavoro.", wrap=True)
        ws.row_dimensions[r + 3].height = 28
        put(ws, f"B{r + 5}", "Domenico Perroni", color=MUTED)
        cell = put(ws, f"C{r + 5}", REPO, color="1F5FA8")
        cell.hyperlink = REPO


def build(o: dict | None = None, path: Path = OUTPUT) -> Foglio:
    o = controlli.run() if o is None else o
    f = Foglio(o)
    f.build(path)
    return f


def main() -> None:
    f = build()
    print(f"foglio: {len(f.checks)} confronti con Python -> {OUTPUT.relative_to(C.ROOT)}")


if __name__ == "__main__":
    main()
