"""Il foglio di calcolo dà gli stessi risultati di Python, formula per formula."""
import pytest

formulas = pytest.importorskip("formulas")

import build_foglio  # noqa: E402
import check_workbook  # noqa: E402
from deterministic import same_content  # noqa: E402


@pytest.fixture(scope="module")
def foglio(out, tmp_path_factory):
    path = tmp_path_factory.mktemp("f") / "foglio.xlsx"
    return build_foglio.build(out, path), path


@pytest.fixture(scope="module")
def valori(foglio):
    """Il foglio calcolato una volta sola: il calcolo in Python puro richiede un minuto."""
    return check_workbook.values_from_formulas(foglio[1])


def test_ogni_formula_torna_con_python(foglio, valori, capsys):
    f, _ = foglio
    assert check_workbook.check(valori, f.wb.sheetnames, len(f.checks)) == 0, capsys.readouterr().out


def test_un_valore_sbagliato_viene_visto(foglio, tmp_path):
    from openpyxl import load_workbook
    f, path = foglio
    wb = load_workbook(path)
    ws = wb["Riconciliazione"]
    row = next(r for r in range(7, 7 + len(f.checks)) if isinstance(ws[f"D{r}"].value, float))
    ws[f"D{row}"].value += 0.01
    broken = tmp_path / "rotto.xlsx"
    wb.save(broken)
    assert check_workbook.check(check_workbook.values_from_formulas(broken), wb.sheetnames, len(f.checks)) == 1


def test_il_foglio_e_riproducibile(out, tmp_path):
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    build_foglio.build(out, a)
    build_foglio.build(out, b)
    assert same_content(a.read_bytes(), b.read_bytes())


def test_un_cedolino_incollato_viene_controllato(foglio, valori, tmp_path):
    """Chi usa il foglio incolla i propri cedolini: una riga nuova, con il minimo sbagliato, deve
    comparire negli esiti, nel riepilogo e nella tranche di novembre senza toccare le formule."""
    from datetime import date
    from openpyxl import load_workbook
    f, path = foglio
    first, last = f.rows
    wb = load_workbook(path)
    ws = wb["Mese"]
    r = next(i for i in range(first, last + 1) if ws[f"A{i}"].value is None)
    riga = ["Z999", "B", "RSSMRA85T10A562S", "4", date(2024, 1, 15), 1.0, "", True,
            1746.68, 0, 0.0, 0.0, 0.0, 1746.68, 160.55, 1586.13, 200.0, 23.79, 1362.34, 150.95, 2.1667, 0.0, 2.1667]
    for j, v in enumerate(riga, start=1):
        ws.cell(row=r, column=j, value=v)
    nuovo = tmp_path / "incollato.xlsx"
    wb.save(nuovo)
    v = check_workbook.values_from_formulas(nuovo)
    S = lambda ref: v[("MESE", ref)]
    assert S(f"AK{r}") == "ERRORE"                       # C01: tabellare fermo a marzo 2025
    assert S(f"AT{r}") == 1
    assert v[("NOVEMBRE 2026", f"A{r}")] == "Z999" and v[("NOVEMBRE 2026", f"G{r}")] == 35.0
    assert v[("RIEPILOGO", "D5")] == valori[("RIEPILOGO", "D5")] + 1    # C01, cliente B
    vuota = r + 1
    assert S(f"AK{vuota}") in ("", None) and S(f"AT{vuota}") in ("", None)
